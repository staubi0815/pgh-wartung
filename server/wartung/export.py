"""Export (docs/05 Abschnitt 3): Foxtag-Importformat und Vollexport, jeweils als ZIP.

Foxtag-Export: je Datenart eine Excel-Datei genau in den Spalten der Foxtag-Vorlagen, nummeriert in Import-Reihenfolge.
Die Spalten kommen aus denselben Definitionen wie der Import (excel_import.ARTEN) – Export und Import passen so
immer zusammen. Exportiert wird nur der gültige Bestand (nicht als gelöscht markiert, Melder nur verbaut).

Vollexport: alle Tabellen vollständig (auch gelöschte und ersetzte Datensätze) als CSV und JSON, dazu das Schema.
Ausgenommen sind nur Anmeldegeheimnisse und das Anmelde-Sicherheitsprotokoll (AUSGENOMMEN unten).

Beide Exporte lesen in einer Lesetransaktion, also einen in sich stimmigen Stand, auch wenn gleichzeitig jemand ändert.
"""
import csv
import io
import json
import re
import zipfile
from contextlib import contextmanager
from datetime import date, datetime

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from . import anlagenart, db
from .excel_import import ARTEN

DATUMSFORMAT = "DD.MM.YYYY"

# Tabellen, die der Vollexport bewusst weglässt, mit Begründung (steht auch in LIESMICH.txt)
AUSGENOMMEN_TABELLEN = {
    "anmeldeversuch": "Sicherheitsprotokoll der Anmeldungen (keine Geschäftsdaten, wird nicht übernommen)",
}
# Spalten mit Anmeldegeheimnissen; zusätzlich gilt: jede Spalte, die auf „_hash“ endet, wird weggelassen
AUSGENOMMEN_SPALTEN = {
    "nutzer": ("einladung_bis", "sitzung_zaehler"),
}


@contextmanager
def _lesestand(con):
    """Lesetransaktion: alle Abfragen sehen denselben Stand (WAL), ohne andere beim Schreiben zu blockieren."""
    if con.in_transaction:
        yield
        return
    con.execute("BEGIN")
    try:
        yield
    finally:
        con.execute("ROLLBACK")


def _zip(dateien):
    """dateien: [(Name, Bytes)] -> ZIP als Bytes."""
    puffer = io.BytesIO()
    zeit = datetime.now(db.TZ).timetuple()[:6]
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
        for name, inhalt in dateien:
            info = zipfile.ZipInfo(name, date_time=zeit)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, inhalt)
    return puffer.getvalue()


def dateiname_sicher(text):
    """Teil eines Dateinamens: nur Buchstaben, Ziffern, Punkt, Minus, Unterstrich."""
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("._") or "ohne_nummer"


# ---------- Foxtag-Format ----------

def _datum(text):
    return date.fromisoformat(text) if text else None


def _mappe(importart, zeilen, blatt):
    """Excel-Datei wie die Foxtag-Vorlage: Kopfzeile (Pflichtspalten mit *), dann je Datensatz eine Zeile."""
    mappe = openpyxl.Workbook()
    ws = mappe.active
    ws.title = blatt[:31]
    ws.append([f"{s}*" if s in importart.pflicht else s for s in importart.spalten])
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
    for werte in zeilen:
        # Steuerzeichen (außer Tab/Zeilenumbruch) sind in Excel-Dateien nicht erlaubt und werden weggelassen
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", w) if isinstance(w, str) else w
                   for w in (werte.get(s) for s in importart.spalten)])
        for zelle in ws[ws.max_row]:
            if isinstance(zelle.value, date):
                zelle.number_format = DATUMSFORMAT
            elif isinstance(zelle.value, str) and zelle.data_type == "f":
                zelle.data_type = "s"   # Text, der mit „=“ beginnt, bleibt Text (keine Formel)
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def _kunden(con):
    return [{"KUNDEN.NUMMER": k["nummer"], "KUNDE.NAME": k["name"], "ADRESSZEILE 1": k["strasse"],
             "ADRESSZEILE 2": k["zusatz"], "PLZ": k["plz"], "ORT": k["ort"], "LAND": k["land"],
             "NOTIZ": k["notiz_intern"]}
            for k in con.execute("SELECT * FROM kunde WHERE geloescht = 0 ORDER BY nummer COLLATE NOCASE")]


def _kontakte(con):
    zeilen = []
    for k in con.execute("SELECT c.*, k.nummer AS kunde_nummer FROM kontakt c LEFT JOIN kunde k ON k.id = c.kunde_id "
                         "WHERE c.geloescht = 0 AND (c.kunde_id IS NULL OR k.geloescht = 0) "
                         "ORDER BY k.nummer COLLATE NOCASE, c.name COLLATE NOCASE"):
        notiz = "\n".join(x for x in (f"Funktion: {k['funktion']}" if k["funktion"] else "", k["notiz"]) if x)
        zeilen.append({"NAME": k["name"], "FIRMA": k["firma"], "EMAIL": k["email"], "TELEFON": k["telefon"],
                       "MOBIL": k["mobil"], "FAX": k["fax"], "NOTIZ": notiz, "KUNDE": k["kunde_nummer"] or ""})
    return zeilen


def _objekte(con):
    sql = ("SELECT o.nummer, o.bezeichnung, k.nummer AS kunde_nummer, "
           " CASE WHEN o.adresse_wie_kunde = 1 THEN k.strasse ELSE o.strasse END AS strasse, "
           " CASE WHEN o.adresse_wie_kunde = 1 THEN k.plz ELSE o.plz END AS plz, "
           " CASE WHEN o.adresse_wie_kunde = 1 THEN k.ort ELSE o.ort END AS ort, "
           " CASE WHEN o.adresse_wie_kunde = 1 THEN k.land ELSE o.land END AS land "
           "FROM objekt o JOIN kunde k ON k.id = o.kunde_id WHERE o.geloescht = 0 AND k.geloescht = 0 "
           "ORDER BY o.nummer COLLATE NOCASE")
    return [{"KUNDE.NUMMER": o["kunde_nummer"], "OBJEKT.NAME": o["bezeichnung"], "OBJEKT.NUMMER": o["nummer"],
             "ADRESSZEILE 1": o["strasse"], "ADRESSZEILE 2": "", "PLZ": o["plz"], "ORT": o["ort"], "LAND": o["land"]}
            for o in con.execute(sql)]


def wartungsanwendung(art):
    """Foxtag-Nummer der Wartungsanwendung: erster Importname der Anlagenart (Rauchwarnmelder: „RWM“)."""
    return art.import_namen[0] if art.import_namen else art.name


def _anlagen(con):
    sql = ("SELECT a.*, o.nummer AS objekt_nummer FROM anlage a JOIN objekt o ON o.id = a.objekt_id "
           "JOIN kunde k ON k.id = o.kunde_id WHERE a.geloescht = 0 AND o.geloescht = 0 AND k.geloescht = 0 "
           "ORDER BY a.nummer COLLATE NOCASE")
    zeilen, liste = [], []
    for a in con.execute(sql).fetchall():
        zeilen.append({"OBJEKT.NUMMER": a["objekt_nummer"],
                       "WARTUNGSANWENDUNG.NUMMER": wartungsanwendung(anlagenart.holen(a["anlagenart"])),
                       "ANLAGE.NUMMER": a["nummer"], "ANLAGE.NAME": a["bezeichnung"], "TECHNIKER.NUMMER": ""})
        liste.append(a)
    return zeilen, liste


def _typen(con, art_schluessel):
    kategorie = {"komponente": "Komponente", "sub_komponente": "Sub-Komponente"}
    return [{"TYP.NAME": t["bezeichnung"], "TYP.HERSTELLER": t["hersteller"], "TYP.MODELL": t["modell"],
             "TYP.KATEGORIE": kategorie[t["kategorie"]], "TYP.LINK": t["datenblatt_link"]}
            for t in con.execute("SELECT * FROM komponententyp WHERE anlagenart = ? AND geloescht = 0 "
                                 "ORDER BY hersteller COLLATE NOCASE, bezeichnung COLLATE NOCASE", (art_schluessel,))]


def gruppenname(bewohner, bezeichnung):
    """Foxtag kennt nur GRUPPE.NAME: „Bewohner, Lage“ (wie in der Foxtag-Vorlage), ohne Bewohner nur die Lage."""
    return f"{bewohner}, {bezeichnung}" if bewohner else bezeichnung


def _komponenten(con, anlage_id):
    sql = ("SELECT k.*, g.nummer AS gruppe_nummer, g.bezeichnung AS gruppe_bezeichnung, g.bewohner, "
           " t.bezeichnung AS typ_name, t.hersteller, t.modell "
           "FROM komponente k JOIN gruppe g ON g.id = k.gruppe_id JOIN komponententyp t ON t.id = k.komponententyp_id "
           "WHERE k.anlage_id = ? AND k.geloescht = 0 AND k.status = 'verbaut' AND g.geloescht = 0 "
           "ORDER BY g.nummer, k.nummer, k.sub_nummer")
    return [{"GRUPPE.NUMMER": k["gruppe_nummer"], "GRUPPE.NAME": gruppenname(k["bewohner"], k["gruppe_bezeichnung"]),
             "NUMMER": k["nummer"], "SUB-NUMMER": k["sub_nummer"], "TYP.NAME": k["typ_name"],
             "TYP.HERSTELLER": k["hersteller"], "TYP.MODELL": k["modell"], "STANDORT": k["raum"],
             "SERIENNUMMER": k["seriennummer"], "QR-CODE": k["barcode"] or "", "BAUJAHR": k["baujahr"],
             "LABEL": None, "LABEL2": None, "LETZTE PRÜFUNG": _datum(k["letzte_pruefung_am"]),
             "INBETRIEBNAHME AM": _datum(k["inbetriebnahme_am"])}
            for k in con.execute(sql, (anlage_id,))]


FOXTAG_LIESMICH = """Export im Foxtag-Importformat (PGH-Wartung), erstellt am {zeit}

Die Dateien in der Reihenfolge ihrer Nummern über den Datenimport von Foxtag einlesen:
{dateien}

Vor dem Import in Foxtag:
- Die Wartungsanwendung für Rauchwarnmelder muss in Foxtag die Nummer „{rwm}“ haben (Spalte
  WARTUNGSANWENDUNG.NUMMER in der Anlagen-Datei), sonst die Spalte vorher anpassen.
- Melder je Anlage über die Anlage in Foxtag importieren (Datei 06_..._<Anlagennummer>.xlsx).

Was nicht oder anders übertragen wird:
- Nur der gültige Bestand: als gelöscht markierte Datensätze fehlen, Melder nur, wenn sie verbaut sind.
- Objekte: Anschrift ist die wirksame Anschrift (bei „wie Kunde“ die des Kunden).
- Kontakte: die Funktion (z. B. Hausmeister) steht in der Notiz, da Foxtag dafür keine Spalte hat.
- Wohnungen: Bewohner und Lage stehen zusammen in GRUPPE.NAME („Bewohner, Lage“).
- Foxtag übernimmt je Melder nur das Datum der letzten Prüfung, nicht den Prüfverlauf; der Verlauf bleibt im
  Vollexport und in den Berichten.
- Noch nicht enthalten (gibt es in PGH-Wartung noch nicht): Labels, Techniker je Anlage, Artikel, Aufträge.

Achtung: Die Dateien enthalten Kunden- und Mieterdaten. Nur auf dem NAS ablegen, nicht per Mail versenden;
an einen anderen Dienst erst nach Abschluss eines Auftragsverarbeitungsvertrags übergeben.
"""


def foxtag(con):
    """Foxtag-Export als ZIP. Rückgabe: (ZIP-Bytes, Anzahl je Datei {Dateiname: Zeilen})."""
    dateien, anzahl = [], {}

    def dazu(name, importart, zeilen, blatt):
        stamm, nr = name.removesuffix(".xlsx"), 2
        while name.lower() in (n.lower() for n in anzahl):  # z. B. „ANL/1“ und „ANL_1“ ergeben denselben Namen
            name, nr = f"{stamm}_{nr}.xlsx", nr + 1
        dateien.append((name, _mappe(importart, zeilen, blatt)))
        anzahl[name] = len(zeilen)

    with _lesestand(con):
        dazu("01_Kunden.xlsx", ARTEN["kunden"][0], _kunden(con), "Kunden")
        dazu("02_Kontakte.xlsx", ARTEN["kontakte"][0], _kontakte(con), "Kontakte")
        dazu("03_Objekte.xlsx", ARTEN["objekte"][0], _objekte(con), "Objekte")
        zeilen, anlagen_liste = _anlagen(con)
        dazu("04_Anlagen.xlsx", ARTEN["anlagen"][0], zeilen, "Anlagen")
        for art in anlagenart.alle().values():
            zeilen = _typen(con, art.schluessel)
            if zeilen:
                dazu(f"05_Typen_{dateiname_sicher(art.name)}.xlsx", ARTEN["typen"][0], zeilen, art.komponente_mehrzahl)
        for a in anlagen_liste:
            art = anlagenart.holen(a["anlagenart"])
            zeilen = _komponenten(con, a["id"])
            if zeilen:
                dazu(f"06_{dateiname_sicher(art.komponente_mehrzahl)}_{dateiname_sicher(a['nummer'])}.xlsx",
                     ARTEN["komponenten"][0], zeilen, art.name)
    liste = "\n".join(f"  {name}  ({n} {'Zeile' if n == 1 else 'Zeilen'})" for name, n in anzahl.items())
    rwm = wartungsanwendung(anlagenart.holen("rauchwarnmelder"))
    text = FOXTAG_LIESMICH.format(zeit=db.jetzt(), dateien=liste, rwm=rwm)
    return _zip([("LIESMICH.txt", text.encode("utf-8"))] + dateien), anzahl


# ---------- Vollexport ----------

def _csv_zelle(wert):
    """Wert für CSV. Text, den Tabellenprogramme als Formel ausführen würden, bekommt ein ' davor
    (Telefonnummern wie „+49 …“ bleiben unverändert). JSON enthält immer den unveränderten Wert."""
    if wert is None:
        return ""
    if isinstance(wert, str) and wert[:1] in ("=", "@", "+", "-", "\t", "\r"):
        if not re.fullmatch(r"[+-][\d\s/().-]*", wert):
            return "'" + wert
    return wert


def tabellen(con):
    """(Tabelle, [exportierte Spalten], Sortierung) aller Tabellen außer den ausgenommenen, alphabetisch.

    Sortierung: Reihenfolge des Anlegens (rowid); Tabellen WITHOUT ROWID nach ihrem Primärschlüssel.
    """
    ergebnis = []
    for t in con.execute("SELECT name, sql FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%' "
                         "ORDER BY name").fetchall():
        name = t["name"]
        if name in AUSGENOMMEN_TABELLEN:
            continue
        info = con.execute(f"PRAGMA table_info({name})").fetchall()
        spalten = [s["name"] for s in info
                   if not s["name"].endswith("_hash") and s["name"] not in AUSGENOMMEN_SPALTEN.get(name, ())]
        if "WITHOUT ROWID" in t["sql"].upper():
            sortierung = ", ".join(s["name"] for s in sorted(info, key=lambda s: s["pk"]) if s["pk"])
        else:
            sortierung = "rowid"
        ergebnis.append((name, spalten, sortierung))
    return ergebnis


VOLL_LIESMICH = """Vollexport PGH-Wartung, erstellt am {zeit}

Inhalt:
- daten.json      alle Tabellen mit allen Datensätzen (maßgeblich; Werte unverändert, leere Werte als null)
- tabellen/*.csv  dieselben Daten je Tabelle (UTF-8, Semikolon, erste Zeile = Spaltennamen; leere Werte = leer;
                  Text, der in Tabellenprogrammen als Formel liefe, beginnt mit ' )
- schema.sql      Aufbau der Datenbank (Tabellen, Prüfregeln)

Enthalten ist der vollständige Stand inklusive als gelöscht markierter (geloescht = 1) und ersetzter Datensätze sowie
des Änderungsprotokolls. Verknüpfungen laufen über die Spalte id (UUID).

Tabellen und Anzahl Datensätze:
{tabellen}

Bewusst nicht enthalten:
{ausgenommen}

Dateien (Fotos, Unterschriften, Berichte) gibt es noch nicht; sie kommen mit den Aufträgen in den Ordner dateien/.

Achtung: Der Export enthält Kunden- und Mieterdaten. Nur auf dem NAS ablegen, nicht per Mail versenden;
an einen anderen Dienst erst nach Abschluss eines Auftragsverarbeitungsvertrags übergeben.
"""


def voll(con):
    """Vollexport als ZIP. Rückgabe: (ZIP-Bytes, Anzahl Datensätze je Tabelle)."""
    daten, dateien, anzahl, schema = {}, [], {}, []
    with _lesestand(con):
        for name, spalten, sortierung in tabellen(con):
            liste = ", ".join(spalten)
            zeilen = [dict(zip(spalten, z)) for z in con.execute(f"SELECT {liste} FROM {name} ORDER BY {sortierung}")]
            daten[name] = zeilen
            anzahl[name] = len(zeilen)
            puffer = io.StringIO()
            schreiber = csv.writer(puffer, delimiter=";", lineterminator="\r\n")
            schreiber.writerow(spalten)
            schreiber.writerows([_csv_zelle(z[s]) for s in spalten] for z in zeilen)
            dateien.append((f"tabellen/{name}.csv", puffer.getvalue().encode("utf-8-sig")))
            schema.append(con.execute("SELECT sql FROM sqlite_master WHERE type = 'table' AND name = ?",
                                      (name,)).fetchone()["sql"] + ";")
    zeit = db.jetzt()
    json_inhalt = json.dumps({"erstellt_am": zeit, "tabellen": daten}, ensure_ascii=False, indent=1)
    ausgenommen = [f"- Tabelle {t}: {grund}" for t, grund in AUSGENOMMEN_TABELLEN.items()]
    ausgenommen.append("- Anmeldegeheimnisse der Nutzer (Spalten *_hash, "
                       + ", ".join(s for liste in AUSGENOMMEN_SPALTEN.values() for s in liste) + ")")
    text = VOLL_LIESMICH.format(zeit=zeit, tabellen="\n".join(f"  {t}: {n}" for t, n in anzahl.items()),
                                ausgenommen="\n".join(ausgenommen))
    return _zip([("LIESMICH.txt", text.encode("utf-8")), ("daten.json", json_inhalt.encode("utf-8")),
                 ("schema.sql", "\n\n".join(schema).encode("utf-8"))] + dateien), anzahl
