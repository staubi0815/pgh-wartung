"""Schritt 11: Foxtag-Export/-Import mit Standard-Techniker (TECHNIKER.NUMMER der Anlage) und Zulassungsnummer."""
import io
import zipfile

import openpyxl
import pytest

from wartung import anlagen, anlagenart, auth, excel_import as imp, export, gruppen, komponenten, kunden, objekte, typen
from wartung.app import erzeuge_app

from hilfen import nutzer_mit_passwort, rolle_id

RWM = anlagenart.holen("rauchwarnmelder")
ANLAGEN_KOPF = ["OBJEKT.NUMMER*", "WARTUNGSANWENDUNG.NUMMER*", "ANLAGE.NUMMER*", "ANLAGE.NAME", "TECHNIKER.NUMMER"]


def entpacken(inhalt):
    with zipfile.ZipFile(io.BytesIO(inhalt)) as z:
        return {n: z.read(n) for n in z.namelist()}


def blatt(inhalt):
    zeilen = list(openpyxl.load_workbook(io.BytesIO(inhalt)).active.iter_rows(values_only=True))
    return list(zeilen[0]), [list(z) for z in zeilen[1:]]


def xlsx(kopf, *zeilen):
    mappe = openpyxl.Workbook()
    mappe.active.append(kopf)
    for z in zeilen:
        mappe.active.append(list(z))
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


@pytest.fixture
def lage(umgebung):
    _, con = umgebung
    k = kunden.anlegen(con, {"art": "hausverwaltung", "name": "Muster", "strasse": "Weg 1", "plz": "90402",
                              "ort": "Nürnberg", "land": "DE"}, None)
    o = objekte.anlegen(con, k, {"nummer": "O-1", "bezeichnung": "Haus", "adresse_wie_kunde": "1"}, None)
    tom = nutzer_mit_passwort(con, "Tom Techniker", "tom@example.org", [rolle_id(con, "techniker")])
    con.execute("UPDATE nutzer SET personalnummer = '17' WHERE id = ?", (tom,))
    typ = typen.anlegen(con, {"anlagenart": "rauchwarnmelder", "bezeichnung": "T", "kategorie": "komponente",
                              "funk": "keine", "batterie": "fest_10j", "aktiv": "1"}, None)

    def anlage(nummer, stamm=None):
        a = anlagen.anlegen(con, o, {"nummer": nummer, "anlagenart": "rauchwarnmelder", "verfahren": "A",
                                     "stammtechniker_id": stamm or ""}, None)
        g = gruppen.holen(con, gruppen.anlegen(con, RWM, a, {"nummer": "1", "zugang": "frei"}, None))
        return a, g
    return {"con": con, "tom": tom, "typ": typ, "anlage": anlage}


def melder(w, g, nummer, **form):
    return komponenten.anlegen(w["con"], RWM, g, {"nummer": str(nummer), "komponententyp_id": w["typ"],
                                                  "baujahr": "2022", **form}, None)


def test_export_anlage_mit_stammtechniker(lage):
    lage["anlage"]("A-1", lage["tom"])
    lage["anlage"]("A-2")
    kopf, zeilen = blatt(entpacken(export.foxtag(lage["con"])[0])["04_Anlagen.xlsx"])
    assert kopf == ANLAGEN_KOPF
    assert [(z[2], z[4]) for z in zeilen] == [("A-1", "17"), ("A-2", None)]


def test_export_stammtechniker_ohne_personalnummer_warnt(lage):
    lage["con"].execute("UPDATE nutzer SET personalnummer = '' WHERE id = ?", (lage["tom"],))
    lage["anlage"]("A-1", lage["tom"])
    liesmich = entpacken(export.foxtag(lage["con"])[0])["LIESMICH.txt"].decode()
    assert "Stammtechniker keine Personalnummer" in liesmich


def test_export_zulassungsnummer_nur_mit_werten(lage):
    a1, g1 = lage["anlage"]("A-1")
    a2, g2 = lage["anlage"]("A-2")
    melder(lage, g1, 1, zulassungsnummer="Z-19-1234")
    melder(lage, g1, 2)
    melder(lage, g2, 1)
    dateien = entpacken(export.foxtag(lage["con"])[0])
    kopf1, z1 = blatt(dateien["06_Komponenten_A-1.xlsx"])
    kopf2, _ = blatt(dateien["06_Komponenten_A-2.xlsx"])
    assert kopf1[-1] == "ZULASSUNGSNUMMER" and kopf2[-1] == "INBETRIEBNAHME AM"  # RWM-Vorlage bleibt unverändert
    assert [z[-1] for z in z1] == ["Z-19-1234", None]
    liesmich = dateien["LIESMICH.txt"].decode()
    assert "06_Komponenten_A-1.xlsx" in liesmich.split("HINWEIS:")[1] and "06_Komponenten_A-2.xlsx" not in \
        liesmich.split("HINWEIS:")[1].split("\n")[0]


def test_rundweg_stammtechniker_und_zulassung(lage, tmp_path):
    a, g = lage["anlage"]("A-1", lage["tom"])
    melder(lage, g, 1, zulassungsnummer="Z-19-1234")
    dateien = entpacken(export.foxtag(lage["con"])[0])
    ziel = erzeuge_app(tmp_path / "ziel").state.con
    tom = nutzer_mit_passwort(ziel, "Tom Techniker", "tom@example.org", [rolle_id(ziel, "techniker")])
    ziel.execute("UPDATE nutzer SET personalnummer = '17' WHERE id = ?", (tom,))
    for name, art in (("01_Kunden.xlsx", "kunden"), ("03_Objekte.xlsx", "objekte"), ("04_Anlagen.xlsx", "anlagen")):
        assert imp.uebernehmen(ziel, art, dateien[name], None).gespeichert
    anlage = anlagen.holen(ziel, ziel.execute("SELECT id FROM anlage").fetchone()["id"])
    assert anlage["stammtechniker_id"] == tom
    assert imp.uebernehmen(ziel, "typen", dateien["05_Komponenten-Typen_Rauchwarnmelder.xlsx"], None,
                           {"anlagenart": "rauchwarnmelder"}).gespeichert
    assert imp.uebernehmen(ziel, "komponenten", dateien["06_Komponenten_A-1.xlsx"], None, {"anlage": anlage}).gespeichert
    assert ziel.execute("SELECT zulassungsnummer FROM komponente").fetchone()[0] == "Z-19-1234"


def test_import_anlage_unbekannter_techniker_nur_hinweis(lage):
    con = lage["con"]
    buero = nutzer_mit_passwort(con, "Bea Büro", "bea@example.org", [rolle_id(con, "buero")])
    con.execute("UPDATE nutzer SET personalnummer = '18' WHERE id = ?", (buero,))
    datei = xlsx(ANLAGEN_KOPF, ["O-1", "RWM", "N-1", "x", "99"], ["O-1", "RWM", "N-2", "y", "18"],
                 ["O-1", "RWM", "N-3", "z", "17"], ["O-1", "RWM", "N-4", "w", None])
    e = imp.uebernehmen(con, "anlagen", datei, None)
    assert e.gespeichert and e.anzahl("neu") == 4
    stamm = dict(con.execute("SELECT nummer, stammtechniker_id FROM anlage").fetchall())
    assert stamm == {"N-1": None, "N-2": None, "N-3": lage["tom"], "N-4": None}  # Büro ist kein Techniker
    meldungen = [m for z in e.zeilen for m in z.meldungen]
    assert sum("Stammtechniker nicht gesetzt" in m for m in meldungen) == 2
