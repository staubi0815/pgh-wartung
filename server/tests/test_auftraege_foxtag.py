"""Aufträge im Foxtag-Format: Export (Kopf wie Vorlage, Werte, Hinweise) und Import (Rundweg, Fehler)."""
import io
import zipfile
from datetime import date, datetime, timedelta

import openpyxl
import pytest

from wartung import anlagen, auftraege, auth, export, kunden, objekte
from wartung import excel_import as imp
from wartung.app import erzeuge_app

from hilfen import anmelden, rolle_id
from test_import import xlsx

FOXTAG_KOPF = ["AUFTRAG.NUMMER", "ANLAGE.NUMMER*", "DATUM*", "AUFTRAGSTYP.NUMMER*", "TECHNIKER.NUMMER*",
               "TECHNIKER.NUMMER2", "TECHNIKER.NUMMER3", "AUFTRAG.HINWEISE"]
MORGEN = date.today() + timedelta(days=1)


def techniker(con, name, email, personalnummer=None):
    return auth.nutzer_anlegen(con, name, email, [rolle_id(con, "techniker")], None, personalnummer=personalnummer)


def auftragsdatei(inhalt):
    with zipfile.ZipFile(io.BytesIO(inhalt)) as z:
        alle = {n: z.read(n) for n in z.namelist()}
    return alle["07_Auftraege.xlsx"], alle["LIESMICH.txt"].decode("utf-8"), alle


@pytest.fixture
def quelle(umgebung):
    _, con = umgebung
    k = kunden.anlegen(con, {"nummer": "K-1", "art": "hausverwaltung", "name": "Muster", "strasse": "Weg 1",
                             "plz": "90402", "ort": "Nürnberg"}, None)
    o = objekte.anlegen(con, k, {"nummer": "O-1", "bezeichnung": "Haus", "adresse_wie_kunde": "1"}, None)
    a = anlagen.holen(con, anlagen.anlegen(con, o, {"nummer": "ANL-1", "anlagenart": "rauchwarnmelder",
                                                    "verfahren": "A"}, None))
    tom = techniker(con, "Tom", "tom@example.org", "P-1")
    ute = techniker(con, "Ute", "ute@example.org")  # ohne Personalnummer

    def plan(**f):
        return auftraege.anlegen(con, a, {"auftragsart": "wartung", "datum": MORGEN.isoformat(), **f}, None)
    ids = {"tom": plan(uhrzeit="08:30", techniker=[tom], hinweise="Schlüssel beim Hausmeister"),
           "pool": plan(auftragsart="installation"),
           "ute": plan(techniker=[ute]),
           "fertig": plan(techniker=[tom])}
    auftraege.status_setzen(con, ids["fertig"], "abgeschlossen", None)
    return con


def test_export_auftraege(quelle):
    datei, liesmich, _ = auftragsdatei(export.foxtag(quelle)[0])
    ws = openpyxl.load_workbook(io.BytesIO(datei)).active
    zeilen = list(ws.iter_rows(values_only=True))
    assert list(zeilen[0]) == FOXTAG_KOPF
    assert [list(z) for z in zeilen[1:]] == [
        ["A-1002", "ANL-1", datetime(MORGEN.year, MORGEN.month, MORGEN.day), "INSTALLATION", None, None, None, None],
        ["A-1003", "ANL-1", datetime(MORGEN.year, MORGEN.month, MORGEN.day), "WARTUNG", None, None, None, None],
        ["A-1001", "ANL-1", datetime(MORGEN.year, MORGEN.month, MORGEN.day, 8, 30), "WARTUNG", "P-1", None, None,
         "Schlüssel beim Hausmeister"],
    ]  # abgeschlossener Auftrag A-1004 fehlt; ganztägige vor solchen mit Uhrzeit
    assert ws["C4"].number_format == export.DATUM_ZEIT_FORMAT and ws["C2"].number_format == export.DATUMSFORMAT
    assert ("    INSTALLATION  (bei uns: Montage / Erstausstattung)\n"
            "    WARTUNG  (bei uns: Wartung / Inspektion)") in liesmich
    assert "NACHTERMIN" not in liesmich  # nur verwendete Auftragstypen
    assert "1 Auftrag/Aufträge ohne Techniker (Pool)" in liesmich
    assert "bei 1 Auftrag/Aufträgen fehlt einem Techniker die Personalnummer" in liesmich


def test_rundweg(quelle, tmp_path):
    """Kunden, Objekte, Anlagen und Aufträge aus dem Export in eine leere Datenbank einlesen."""
    _, _, dateien = auftragsdatei(export.foxtag(quelle)[0])
    ziel = erzeuge_app(tmp_path / "ziel").state.con
    techniker(ziel, "Tom", "tom@example.org", "P-1")
    for name, art in (("01_Kunden.xlsx", "kunden"), ("03_Objekte.xlsx", "objekte"), ("04_Anlagen.xlsx", "anlagen")):
        assert imp.uebernehmen(ziel, art, dateien[name], None).gespeichert
    e = imp.uebernehmen(ziel, "auftraege", dateien["07_Auftraege.xlsx"], None)
    assert e.gespeichert and [z.status for z in e.zeilen] == ["neu"] * 3, [z.meldungen for z in e.zeilen]
    assert e.zeilen[0].meldungen == ["Ohne Techniker: Pool-Auftrag."]
    u = auftraege.holen(ziel, ziel.execute("SELECT id FROM auftrag WHERE nummer = 'A-1001'").fetchone()[0])
    assert (u["datum"], u["uhrzeit"], u["auftragsart"], u["status"], u["hinweise"]) == \
        (MORGEN.isoformat(), "08:30", "wartung", "geplant", "Schlüssel beim Hausmeister")
    assert [t["name"] for t in auftraege.techniker(ziel, u["id"])] == ["Tom"]
    assert ziel.execute("SELECT auftragsart FROM auftrag WHERE nummer = 'A-1002'").fetchone()[0] == "installation"
    # zweites Einlesen: alles vorhanden, nichts doppelt
    nochmal = imp.uebernehmen(ziel, "auftraege", dateien["07_Auftraege.xlsx"], None)
    assert [z.status for z in nochmal.zeilen] == ["vorhanden"] * 3
    assert ziel.execute("SELECT COUNT(*) FROM auftrag").fetchone()[0] == 3


def test_import_fehler_und_hinweise(quelle):
    con = quelle
    gestern = date.today() - timedelta(days=1)
    kopf = FOXTAG_KOPF + ["Hilfe für dieses Vorlage"]
    datei = xlsx(kopf,
                 ["JOB-1", "ANL-1", gestern, "Wartung / Inspektion", "P-1", None, None, "nachgetragen", "Hilfe …"],
                 [None, "ANL-9", MORGEN, "WARTUNG", None, None, None, None, None],
                 [None, "ANL-1", MORGEN, "SERVICE-1", None, None, None, None, None],
                 [None, "ANL-1", MORGEN, "WARTUNG", "P-99", None, None, None, None],
                 [None, "ANL-1", None, "WARTUNG", None, None, None, None, None],
                 ["job-1", "ANL-1", MORGEN, "WARTUNG", None, None, None, None, None],
                 [None, None, None, None, None, None, None, None, "Löschen Sie optionale Spalten"])
    e = imp.uebernehmen(con, "auftraege", datei, None)
    assert not e.gespeichert and e.hinweise == []  # Hilfespalte wird nicht als unbekannt gemeldet
    meldungen = {z.nr: (z.status, z.meldungen) for z in e.zeilen}
    assert meldungen[2] == ("neu", ["Datum liegt in der Vergangenheit."])  # Name statt Nummer ist auch erlaubt
    assert "Anlage ANL-9 gibt es nicht" in meldungen[3][1][0]
    assert "SERVICE-1" in meldungen[4][1][0] and "WARTUNG" in meldungen[4][1][0]
    assert "Personalnummer P-99" in meldungen[5][1][0]
    assert meldungen[6][1] == ["Datum fehlt."]
    assert meldungen[7][1] == ["Doppelt in der Datei (wie Zeile 2)."]
    assert 8 not in meldungen  # reine Hilfezeile übersprungen
    # nichts gespeichert, auch die fehlerfreie erste Zeile nicht
    assert con.execute("SELECT COUNT(*) FROM auftrag WHERE nummer = 'JOB-1'").fetchone()[0] == 0


def test_vorlage_auftraege(umgebung):
    c, _ = umgebung
    anmelden(c)
    r = c.get("/verwaltung/import/vorlage/auftraege.xlsx")
    kopf = next(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert list(kopf) == FOXTAG_KOPF
    assert "Aufträge" in c.get("/verwaltung/import").text
